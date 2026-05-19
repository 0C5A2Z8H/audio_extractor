#!/usr/bin/env python3
r"""从服务器小时级录音中提取带标注的卫星音频片段。

服务器录音目录通常如下：

    D:\AudioRecord\LX20260407\LX20260407-010000.m4a

脚本读取 Excel 标注表，使用起始时间、结束时间和类别列定位真实录音覆盖区间，
再调用 ffmpeg 切割并解码为 WAV，同时写出 manifest 以便后续追溯。
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from typing import Iterable


INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
SOURCE_NAME_RE = re.compile(r"^(?P<station>[A-Za-z]+)(?P<day>\d{8})-(?P<clock>\d{6})\.(?P<ext>\w+)$")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="从小时级卫星录音中切割带标注的音频片段。"
    )
    parser.add_argument("--audio-root", type=Path, default=Path.cwd())
    parser.add_argument("--excel", type=Path, required=True)
    parser.add_argument("--password", default=None)
    parser.add_argument("--sheet", default=None, help="Worksheet name. Defaults to active sheet.")
    parser.add_argument("--station", default="LX")
    parser.add_argument("--source-ext", default="m4a")
    parser.add_argument("--start-col", default="F")
    parser.add_argument("--end-col", default="G")
    parser.add_argument("--label-col", default="T")
    parser.add_argument("--invalid-col", default="N")
    parser.add_argument("--invalid-marker", default="无效信号")
    parser.add_argument("--include-invalid-signals", action="store_true")
    parser.add_argument("--first-data-row", type=int, default=2)
    parser.add_argument("--output-root", type=Path, default=None)
    parser.add_argument(
        "--clear-output-root",
        action="store_true",
        help="Delete existing files under --output-root before extraction. This is now the default for real runs.",
    )
    parser.add_argument("--keep-output-root", action="store_true", help="Do not clear --output-root before extraction.")
    parser.add_argument("--manifest", type=Path, default=None)
    parser.add_argument("--ffmpeg", default="ffmpeg")
    parser.add_argument("--ffprobe", default="ffprobe")
    parser.add_argument(
        "--source-match-mode",
        choices=("coverage", "exact-hour"),
        default="coverage",
        help="coverage scans real recording intervals; exact-hour uses STATIONYYYYMMDD-HH0000.ext.",
    )
    parser.add_argument("--sample-rate", type=int, default=None)
    parser.add_argument("--channels", type=int, default=None)
    parser.add_argument("--pre-roll-sec", type=float, default=0.0)
    parser.add_argument("--post-roll-sec", type=float, default=0.0)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--no-progress", action="store_true", help="Disable console progress bar.")
    parser.add_argument(
        "--progress-interval-sec",
        type=float,
        default=0.2,
        help="Minimum seconds between progress bar refreshes.",
    )
    return parser.parse_args()


def load_workbook(path: Path, password: str | None):
    try:
        from openpyxl import load_workbook as openpyxl_load_workbook
    except ImportError as exc:
        raise SystemExit("Missing dependency: install openpyxl on the server.") from exc

    try:
        return openpyxl_load_workbook(path, data_only=True, read_only=True)
    except Exception as first_error:
        if not password:
            raise SystemExit(
                f"Could not open workbook {path}. If it is encrypted, pass --password."
            ) from first_error

        try:
            import msoffcrypto
        except ImportError as exc:
            raise SystemExit(
                "Workbook appears encrypted. Install msoffcrypto-tool or save an "
                "unencrypted copy of the workbook on the server."
            ) from exc

        decrypted = io.BytesIO()
        with path.open("rb") as handle:
            office_file = msoffcrypto.OfficeFile(handle)
            office_file.load_key(password=password)
            office_file.decrypt(decrypted)
        decrypted.seek(0)
        return openpyxl_load_workbook(decrypted, data_only=True, read_only=True)


def column_index(column: str) -> int:
    result = 0
    for char in column.strip().upper():
        if not ("A" <= char <= "Z"):
            raise ValueError(f"Invalid Excel column: {column}")
        result = result * 26 + (ord(char) - ord("A") + 1)
    return result


def column_letter(index: int) -> str:
    letters = []
    while index:
        index, remainder = divmod(index - 1, 26)
        letters.append(chr(ord("A") + remainder))
    return "".join(reversed(letters))


def normalize_cell_text(value) -> str:
    return re.sub(r"\s+", "", str(value or "").strip()).lower()


def row_value(row: tuple[object, ...], index: int | None):
    if index is None or len(row) < index:
        return None
    return row[index - 1]


def parse_excel_datetime(value) -> dt.datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time.min)
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel
        except ImportError as exc:
            raise SystemExit("Missing dependency: install openpyxl on the server.") from exc
        parsed = from_excel(value)
        if isinstance(parsed, dt.datetime):
            return parsed.replace(tzinfo=None)
        if isinstance(parsed, dt.time):
            return dt.datetime.combine(dt.date.today(), parsed)
        return dt.datetime.combine(parsed, dt.time.min)

    text = str(value).strip()
    if not text:
        return None

    formats = (
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y-%m-%d %H:%M",
        "%Y%m%d %H:%M:%S",
        "%Y%m%d-%H%M%S",
        "%Y%m%d-%H%M",
    )
    for fmt in formats:
        try:
            return dt.datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise ValueError(f"Unsupported datetime value: {value!r}")


def sanitize_label(value) -> str:
    label = str(value).strip() if value is not None else ""
    label = INVALID_FILENAME_CHARS.sub("_", label)
    label = re.sub(r"\s+", " ", label).strip(" .")
    return label


def floor_to_hour(value: dt.datetime) -> dt.datetime:
    return value.replace(minute=0, second=0, microsecond=0)


def source_path(audio_root: Path, station: str, ext: str, timestamp: dt.datetime) -> Path:
    day = timestamp.strftime("%Y%m%d")
    hour = timestamp.strftime("%H")
    return audio_root / f"{station}{day}" / f"{station}{day}-{hour}0000.{ext.lstrip('.')}"


def iter_hour_chunks(
    start: dt.datetime, end: dt.datetime
) -> Iterable[tuple[dt.datetime, dt.datetime, Path | None, dt.datetime]]:
    cursor = start
    while cursor < end:
        next_hour = floor_to_hour(cursor) + dt.timedelta(hours=1)
        chunk_end = min(end, next_hour)
        yield cursor, chunk_end, None, floor_to_hour(cursor)
        cursor = chunk_end


def parse_source_start(path: Path, station: str, ext: str) -> dt.datetime | None:
    match = SOURCE_NAME_RE.match(path.name)
    if not match:
        return None
    if match.group("station") != station or match.group("ext").lower() != ext.lstrip(".").lower():
        return None
    return dt.datetime.strptime(match.group("day") + match.group("clock"), "%Y%m%d%H%M%S")


def ffprobe_duration(ffprobe: str, path: Path) -> tuple[float | None, str]:
    command = [
        ffprobe,
        "-v",
        "error",
        "-show_entries",
        "format=duration",
        "-of",
        "json",
        str(path),
    ]
    result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if result.returncode != 0:
        return None, result.stderr.strip()
    try:
        payload = json.loads(result.stdout)
        return float(payload["format"]["duration"]), ""
    except Exception as exc:
        return None, f"Could not parse ffprobe output: {exc}"


def get_day_inventory(args: argparse.Namespace, day: str) -> list[dict[str, object]]:
    if not hasattr(args, "_source_inventory"):
        args._source_inventory = {}
    if day in args._source_inventory:
        return args._source_inventory[day]

    day_dir = args.audio_root / f"{args.station}{day}"
    ext = args.source_ext.lstrip(".")
    intervals: list[dict[str, object]] = []
    if day_dir.exists():
        for path in sorted(day_dir.glob(f"{args.station}{day}-*.{ext}")):
            start = parse_source_start(path, args.station, args.source_ext)
            if start is None:
                continue
            duration, probe_error = ffprobe_duration(args.ffprobe, path)
            if duration is None:
                intervals.append(
                    {
                        "path": path,
                        "start": start,
                        "end": None,
                        "probe_error": probe_error,
                    }
                )
                continue
            intervals.append(
                {
                    "path": path,
                    "start": start,
                    "end": start + dt.timedelta(seconds=duration),
                    "probe_error": "",
                }
            )
    intervals.sort(key=lambda item: item["start"])
    args._source_inventory[day] = intervals
    return intervals


def dates_between(start: dt.datetime, end: dt.datetime) -> list[str]:
    cursor = start.date()
    last = end.date()
    days: list[str] = []
    while cursor <= last:
        days.append(cursor.strftime("%Y%m%d"))
        cursor += dt.timedelta(days=1)
    return days


def resolve_exact_hour_segments(
    args: argparse.Namespace, start: dt.datetime, end: dt.datetime
) -> list[tuple[Path, dt.datetime, dt.datetime, dt.datetime]]:
    segments: list[tuple[Path, dt.datetime, dt.datetime, dt.datetime]] = []
    for chunk_start, chunk_end, _, file_start in iter_hour_chunks(start, end):
        path = source_path(args.audio_root, args.station, args.source_ext, chunk_start)
        if not path.exists():
            raise FileNotFoundError(f"Missing source recording: {path}")
        segments.append((path, chunk_start, chunk_end, file_start))
    return segments


def resolve_coverage_segments(
    args: argparse.Namespace, start: dt.datetime, end: dt.datetime
) -> list[tuple[Path, dt.datetime, dt.datetime, dt.datetime]]:
    intervals: list[dict[str, object]] = []
    for day in dates_between(start - dt.timedelta(days=1), end):
        intervals.extend(get_day_inventory(args, day))
    intervals = [item for item in intervals if item["end"] is not None]
    intervals.sort(key=lambda item: item["start"])

    segments: list[tuple[Path, dt.datetime, dt.datetime, dt.datetime]] = []
    cursor = start
    while cursor < end:
        covering = [
            item
            for item in intervals
            if item["start"] <= cursor and item["end"] and cursor < item["end"]
        ]
        if not covering:
            future = [item for item in intervals if item["start"] > cursor]
            next_hint = ""
            if future:
                next_item = future[0]
                next_hint = f"; next recording starts at {next_item['start']} ({next_item['path']})"
            raise FileNotFoundError(f"No source recording covers {cursor}{next_hint}")

        item = covering[-1]
        file_start = item["start"]
        file_end = item["end"]
        segment_end = min(end, file_end)
        if segment_end <= cursor:
            raise RuntimeError(f"Invalid recording interval for {item['path']}")
        segments.append((item["path"], cursor, segment_end, file_start))
        cursor = segment_end
    return segments


def ffmpeg_cut_command(
    ffmpeg: str,
    input_path: Path,
    output_path: Path,
    offset_sec: float,
    duration_sec: float,
    overwrite: bool,
    sample_rate: int | None,
    channels: int | None,
    seek_mode: str = "output",
    ignore_decode_errors: bool = False,
) -> list[str]:
    command = [ffmpeg, "-hide_banner", "-loglevel", "error"]
    command.append("-y" if overwrite else "-n")
    if ignore_decode_errors:
        command.extend(["-err_detect", "ignore_err", "-fflags", "+discardcorrupt"])
    if seek_mode == "input":
        command.extend(["-ss", f"{offset_sec:.3f}", "-i", str(input_path), "-t", f"{duration_sec:.3f}"])
    else:
        command.extend(["-i", str(input_path), "-ss", f"{offset_sec:.3f}", "-t", f"{duration_sec:.3f}"])
    command.extend(["-vn", "-acodec", "pcm_s16le"])
    if sample_rate:
        command.extend(["-ar", str(sample_rate)])
    if channels:
        command.extend(["-ac", str(channels)])
    command.append(str(output_path))
    return command


def run_command(command: list[str]) -> None:
    result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if result.returncode != 0:
        joined = " ".join(command)
        raise RuntimeError(f"Command failed: {joined}\n{result.stderr.strip()}")


def remove_partial_output(path: Path) -> None:
    if path.exists():
        path.unlink()


def run_cut_with_fallback(
    args: argparse.Namespace,
    input_path: Path,
    output_path: Path,
    offset_sec: float,
    duration_sec: float,
    overwrite: bool,
) -> str:
    primary = ffmpeg_cut_command(
        args.ffmpeg,
        input_path,
        output_path,
        offset_sec,
        duration_sec,
        overwrite,
        args.sample_rate,
        args.channels,
    )
    try:
        run_command(primary)
        return "primary"
    except RuntimeError as primary_error:
        remove_partial_output(output_path)
        fallback = ffmpeg_cut_command(
            args.ffmpeg,
            input_path,
            output_path,
            offset_sec,
            duration_sec,
            True,
            args.sample_rate,
            args.channels,
            seek_mode="input",
            ignore_decode_errors=True,
        )
        try:
            run_command(fallback)
            return "recover"
        except RuntimeError as fallback_error:
            raise RuntimeError(
                f"Primary ffmpeg cut failed, fallback also failed.\n"
                f"Primary error:\n{primary_error}\n\nFallback error:\n{fallback_error}"
            ) from fallback_error


def concat_wavs(ffmpeg: str, parts: list[Path], output_path: Path, overwrite: bool) -> None:
    with tempfile.TemporaryDirectory(prefix="audio_concat_") as temp_dir:
        list_path = Path(temp_dir) / "inputs.txt"
        lines = []
        for part in parts:
            safe_path = part.resolve().as_posix().replace("'", "'\\''")
            lines.append(f"file '{safe_path}'")
        list_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

        command = [
            ffmpeg,
            "-hide_banner",
            "-loglevel",
            "error",
            "-y" if overwrite else "-n",
            "-f",
            "concat",
            "-safe",
            "0",
            "-i",
            str(list_path),
            "-c",
            "copy",
            str(output_path),
        ]
        run_command(command)


def output_filename(
    station: str,
    row_index: int,
    start: dt.datetime,
    end: dt.datetime,
    quality_suffix: str = "",
) -> str:
    suffix = f"_{quality_suffix}" if quality_suffix else ""
    return (
        f"{station}_{start:%Y%m%d_%H%M%S}_{end:%Y%m%d_%H%M%S}"
        f"_row{row_index:05d}{suffix}.wav"
    )


def mark_recovered_output(
    output_dir: Path,
    output_path: Path,
    station: str,
    row_index: int,
    start: dt.datetime,
    end: dt.datetime,
    overwrite: bool,
) -> Path:
    recovered_path = output_dir / output_filename(station, row_index, start, end, "recover")
    if recovered_path == output_path:
        return output_path
    if recovered_path.exists() and not overwrite:
        raise FileExistsError(f"Recovered output exists, use --overwrite to replace: {recovered_path}")
    if recovered_path.exists():
        recovered_path.unlink()
    output_path.replace(recovered_path)
    return recovered_path


def safe_clear_output_root(output_root: Path, audio_root: Path) -> None:
    resolved_output = output_root.resolve()
    resolved_audio = audio_root.resolve()

    if resolved_output.name.lower() != "raw":
        raise SystemExit(
            f"Refusing to clear {resolved_output}: --clear-output-root only accepts a directory named 'raw'."
        )
    if resolved_output.anchor == str(resolved_output):
        raise SystemExit(f"Refusing to clear drive root: {resolved_output}")
    if resolved_output == resolved_audio or resolved_audio in resolved_output.parents:
        raise SystemExit(
            f"Refusing to clear {resolved_output}: output root must not be inside audio root {resolved_audio}."
        )
    if not resolved_output.exists():
        return
    if not resolved_output.is_dir():
        raise SystemExit(f"Refusing to clear non-directory output root: {resolved_output}")

    for child in resolved_output.iterdir():
        if child.is_dir():
            shutil.rmtree(child)
        else:
            child.unlink()


def extract_clip(
    args: argparse.Namespace,
    row_index: int,
    label: str,
    start: dt.datetime,
    end: dt.datetime,
) -> tuple[Path, list[Path], str]:
    output_dir = args.output_root / label
    output_path = output_dir / output_filename(args.station, row_index, start, end)

    if args.source_match_mode == "exact-hour":
        segments = resolve_exact_hour_segments(args, start, end)
    else:
        segments = resolve_coverage_segments(args, start, end)
    sources = [segment[0] for segment in segments]

    if args.dry_run:
        return output_path, sources, "dry_run"

    output_dir.mkdir(parents=True, exist_ok=True)
    if output_path.exists() and not args.overwrite:
        raise FileExistsError(f"Output exists, use --overwrite to replace: {output_path}")

    decode_statuses: list[str] = []
    if len(segments) == 1:
        source, segment_start, segment_end, file_start = segments[0]
        offset_sec = (segment_start - file_start).total_seconds()
        duration_sec = (segment_end - segment_start).total_seconds()
        decode_status = run_cut_with_fallback(
            args,
            source,
            output_path,
            offset_sec,
            duration_sec,
            args.overwrite,
        )
        decode_statuses.append(decode_status)
        if decode_status == "recover":
            output_path = mark_recovered_output(
                output_dir, output_path, args.station, row_index, start, end, args.overwrite
            )
        return output_path, sources, ",".join(decode_statuses)

    with tempfile.TemporaryDirectory(prefix="audio_extract_") as temp_dir:
        temp_root = Path(temp_dir)
        parts: list[Path] = []
        for index, (source, segment_start, segment_end, file_start) in enumerate(segments):
            offset_sec = (segment_start - file_start).total_seconds()
            duration_sec = (segment_end - segment_start).total_seconds()
            part_path = temp_root / f"part_{index:03d}.wav"
            decode_status = run_cut_with_fallback(
                args,
                source,
                part_path,
                offset_sec,
                duration_sec,
                True,
            )
            decode_statuses.append(decode_status)
            parts.append(part_path)
        concat_wavs(args.ffmpeg, parts, output_path, args.overwrite)

    if "recover" in decode_statuses:
        output_path = mark_recovered_output(
            output_dir, output_path, args.station, row_index, start, end, args.overwrite
        )

    return output_path, sources, ",".join(decode_statuses)


def select_sheet(workbook, sheet_name: str | None):
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise SystemExit(
                f"Worksheet {sheet_name!r} not found. Available sheets: {workbook.sheetnames}"
            )
        return workbook[sheet_name]
    return workbook.active


class ProgressReporter:
    def __init__(self, enabled: bool, total: int | None, interval_sec: float) -> None:
        self.enabled = enabled
        self.total = total if total and total > 0 else None
        self.interval_sec = max(interval_sec, 0.0)
        self.last_update = 0.0
        self.started_at = time.monotonic()
        self.last_message_len = 0

    def update(
        self,
        processed: int,
        written: int,
        skipped: int,
        failed: int,
        row_index: int | None = None,
        label: str | None = None,
        force: bool = False,
    ) -> None:
        if not self.enabled:
            return
        now = time.monotonic()
        if not force and now - self.last_update < self.interval_sec:
            return
        self.last_update = now

        elapsed = now - self.started_at
        stats = f"elapsed={format_duration(elapsed)} ok={written} skipped={skipped} failed={failed}"

        if self.total:
            ratio = min(processed / self.total, 1.0)
            width = 30
            filled = int(width * ratio)
            bar = "#" * filled + "-" * (width - filled)
            message = f"\rExtracting [{bar}] {processed}/{self.total} {ratio * 100:5.1f}% {stats}"
        else:
            message = f"\rExtracting processed={processed} {stats}"

        message = message.lstrip("\r")[:180]
        padding = " " * max(self.last_message_len - len(message), 0)
        sys.stdout.write("\r" + message + padding)
        sys.stdout.flush()
        self.last_message_len = len(message)

    def finish(self) -> None:
        if self.enabled:
            sys.stdout.write("\n")
            sys.stdout.flush()


def format_duration(seconds: float) -> str:
    seconds_int = int(seconds)
    hours, remainder = divmod(seconds_int, 3600)
    minutes, secs = divmod(remainder, 60)
    if hours:
        return f"{hours:d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def shorten(value: str, max_length: int) -> str:
    if len(value) <= max_length:
        return value
    return value[: max_length - 3] + "..."


def estimate_total_rows(worksheet, first_data_row: int, limit: int | None) -> int | None:
    max_row = getattr(worksheet, "max_row", None)
    if max_row is None:
        return limit
    total = max(max_row - first_data_row + 1, 0)
    if limit is not None:
        total = min(total, limit)
    return total


def main() -> int:
    args = parse_args()
    args.audio_root = args.audio_root.resolve()
    args.excel = args.excel.resolve()
    if args.output_root:
        args.output_root = args.output_root.resolve()
    else:
        args.output_root = (args.audio_root / "SatelliteAudio" / "raw").resolve()
    if args.manifest:
        args.manifest = args.manifest.resolve()
    else:
        args.manifest = (args.output_root.parent / "extraction_manifest.csv").resolve()

    if not shutil.which(args.ffmpeg):
        raise SystemExit(f"ffmpeg not found: {args.ffmpeg}")
    if args.source_match_mode == "coverage" and not shutil.which(args.ffprobe):
        raise SystemExit(f"ffprobe not found: {args.ffprobe}")

    should_clear_output = (not args.keep_output_root) or args.clear_output_root
    if should_clear_output:
        if args.dry_run:
            print("output-root cleanup skipped because --dry-run is set.")
        else:
            safe_clear_output_root(args.output_root, args.audio_root)

    workbook = load_workbook(args.excel, args.password)
    worksheet = select_sheet(workbook, args.sheet)

    start_idx = column_index(args.start_col)
    end_idx = column_index(args.end_col)
    label_idx = column_index(args.label_col)
    invalid_idx = None if args.invalid_col.lower() in {"", "none"} else column_index(args.invalid_col)
    print(
        "Resolved columns: "
        f"start={column_letter(start_idx)}, end={column_letter(end_idx)}, "
        f"label={column_letter(label_idx)}, "
        f"invalid={column_letter(invalid_idx) if invalid_idx else 'none'}"
    )

    args.manifest.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "row_index",
        "sample_id",
        "label",
        "start_time",
        "end_time",
        "duration_sec",
        "output_path",
        "source_files",
        "invalid_signal",
        "decode_status",
        "status",
        "message",
    ]

    processed = 0
    written = 0
    skipped = 0
    failed = 0
    total_rows = estimate_total_rows(worksheet, args.first_data_row, args.limit)
    progress = ProgressReporter(
        enabled=not args.no_progress,
        total=total_rows,
        interval_sec=args.progress_interval_sec,
    )

    try:
        progress.update(processed, written, skipped, failed, force=True)
        with args.manifest.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()

            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if row_index < args.first_data_row:
                    continue
                if args.limit is not None and processed >= args.limit:
                    break
                processed += 1

                raw_invalid = row_value(row, invalid_idx)
                invalid_signal = normalize_cell_text(raw_invalid) == normalize_cell_text(args.invalid_marker)
                raw_label = row_value(row, label_idx)
                label = sanitize_label(raw_label)
                progress.update(
                    processed,
                    written,
                    skipped,
                    failed,
                    row_index=row_index,
                    label=label,
                    force=True,
                )
                if invalid_signal and not args.include_invalid_signals:
                    skipped += 1
                    writer.writerow(
                        {
                            "row_index": row_index,
                            "sample_id": "",
                            "label": str(raw_label or ""),
                            "start_time": "",
                            "end_time": "",
                            "duration_sec": "",
                            "output_path": "",
                            "source_files": "",
                            "invalid_signal": str(raw_invalid or ""),
                            "decode_status": "",
                            "status": "skipped",
                            "message": "invalid_signal",
                        }
                    )
                    progress.update(processed, written, skipped, failed)
                    continue

                if not label or label in {"?", "？"}:
                    skipped += 1
                    writer.writerow(
                        {
                            "row_index": row_index,
                            "sample_id": "",
                            "label": str(raw_label or ""),
                            "start_time": "",
                            "end_time": "",
                            "duration_sec": "",
                            "output_path": "",
                            "source_files": "",
                            "invalid_signal": str(raw_invalid or ""),
                            "decode_status": "",
                            "status": "skipped",
                            "message": "unknown_or_unresolved_signal",
                        }
                    )
                    progress.update(processed, written, skipped, failed)
                    continue

                try:
                    raw_start = row_value(row, start_idx)
                    raw_end = row_value(row, end_idx)
                    start = parse_excel_datetime(raw_start)
                    end = parse_excel_datetime(raw_end)
                    if start is None or end is None:
                        raise ValueError("Missing start or end time")

                    start = start - dt.timedelta(seconds=args.pre_roll_sec)
                    end = end + dt.timedelta(seconds=args.post_roll_sec)
                    if end <= start:
                        raise ValueError(f"End time must be after start time: {start} -> {end}")

                    output_path, sources, decode_status = extract_clip(args, row_index, label, start, end)
                    duration_sec = (end - start).total_seconds()
                    sample_id = output_path.stem
                    written += 1
                    writer.writerow(
                        {
                            "row_index": row_index,
                            "sample_id": sample_id,
                            "label": label,
                            "start_time": start.isoformat(sep=" "),
                            "end_time": end.isoformat(sep=" "),
                            "duration_sec": f"{duration_sec:.3f}",
                            "output_path": str(output_path),
                            "source_files": ";".join(str(path) for path in sources),
                            "invalid_signal": str(raw_invalid or ""),
                            "decode_status": decode_status,
                            "status": "dry_run" if args.dry_run else "ok",
                            "message": "",
                        }
                    )
                except Exception as exc:
                    failed += 1
                    writer.writerow(
                        {
                            "row_index": row_index,
                            "sample_id": "",
                            "label": label,
                            "start_time": "",
                            "end_time": "",
                            "duration_sec": "",
                            "output_path": "",
                            "source_files": "",
                            "invalid_signal": str(raw_invalid or ""),
                            "decode_status": "",
                            "status": "failed",
                            "message": str(exc),
                        }
                    )
                finally:
                    progress.update(processed, written, skipped, failed)
    finally:
        progress.update(processed, written, skipped, failed, force=True)
        progress.finish()

    print(
        "Extraction complete: "
        f"processed={processed}, written={written}, skipped={skipped}, "
        f"failed={failed}, manifest={args.manifest}"
    )
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
