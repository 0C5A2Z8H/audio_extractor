#!/usr/bin/env python3
"""报告 Excel 标注中未被真实录音完整覆盖的行。"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import re
import sys
import time
from pathlib import Path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="对比 Excel 标注时间段和服务器录音覆盖区间。"
    )
    parser.add_argument("--excel", type=Path, required=True)
    parser.add_argument("--password", default=None)
    parser.add_argument("--sheet", default=None)
    parser.add_argument("--inventory", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--summary", type=Path, default=None)
    parser.add_argument("--start-col", default="F")
    parser.add_argument("--end-col", default="G")
    parser.add_argument("--label-col", default="T")
    parser.add_argument("--invalid-col", default="N")
    parser.add_argument("--invalid-marker", default="无效信号")
    parser.add_argument("--include-invalid-signals", action="store_true")
    parser.add_argument("--first-data-row", type=int, default=2)
    parser.add_argument("--ignore-dates", nargs="*", default=["20260311", "20260312", "20260313"])
    parser.add_argument("--no-progress", action="store_true")
    parser.add_argument("--progress-interval-sec", type=float, default=0.2)
    return parser.parse_args()


def column_index(column: str) -> int:
    result = 0
    for char in column.strip().upper():
        if not ("A" <= char <= "Z"):
            raise ValueError(f"Invalid Excel column: {column}")
        result = result * 26 + (ord(char) - ord("A") + 1)
    return result


def column_letter(index: int | None) -> str:
    if not index:
        return "none"
    letters = []
    while index:
        index, remainder = divmod(index - 1, 26)
        letters.append(chr(ord("A") + remainder))
    return "".join(reversed(letters))


def normalize_cell_text(value) -> str:
    return re.sub(r"\s+", "", str(value or "").strip()).lower()


def row_value(row: tuple[object, ...], index: int | None):
    if index is None or len(row) < index:
        return None
    return row[index - 1]


def load_workbook(path: Path, password: str | None):
    from openpyxl import load_workbook

    try:
        return load_workbook(path, data_only=True, read_only=True)
    except Exception:
        if not password:
            raise
        import msoffcrypto

        decrypted = io.BytesIO()
        with path.open("rb") as handle:
            office_file = msoffcrypto.OfficeFile(handle)
            office_file.load_key(password=password)
            office_file.decrypt(decrypted)
        decrypted.seek(0)
        return load_workbook(decrypted, data_only=True, read_only=True)


def parse_excel_datetime(value) -> dt.datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time.min)
    if isinstance(value, (int, float)):
        from openpyxl.utils.datetime import from_excel

        parsed = from_excel(value)
        if isinstance(parsed, dt.datetime):
            return parsed.replace(tzinfo=None)
        if isinstance(parsed, dt.date):
            return dt.datetime.combine(parsed, dt.time.min)
        return None

    text = str(value).strip()
    if not text:
        return None
    for fmt in (
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y-%m-%d %H:%M",
        "%Y%m%d-%H%M%S",
        "%Y%m%d-%H%M",
    ):
        try:
            return dt.datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise ValueError(f"Unsupported datetime value: {value!r}")


def sanitize_label(value) -> str:
    return str(value).strip() if value is not None else ""


class ProgressReporter:
    def __init__(self, enabled: bool, total: int | None, interval_sec: float) -> None:
        self.enabled = enabled
        self.total = total if total and total > 0 else None
        self.interval_sec = max(interval_sec, 0.0)
        self.last_update = 0.0
        self.started_at = time.monotonic()

    def update(self, current: int, force: bool = False) -> None:
        if not self.enabled:
            return
        now = time.monotonic()
        if not force and now - self.last_update < self.interval_sec:
            return
        self.last_update = now
        elapsed = format_duration(now - self.started_at)
        if self.total:
            ratio = min(current / self.total, 1.0)
            width = 30
            filled = int(width * ratio)
            bar = "#" * filled + "-" * (width - filled)
            message = f"\rCoverage [{bar}] {current}/{self.total} {ratio * 100:5.1f}% elapsed={elapsed}"
        else:
            message = f"\rCoverage processed={current} elapsed={elapsed}"
        sys.stderr.write(message[:180].ljust(180))
        sys.stderr.flush()

    def finish(self) -> None:
        if self.enabled:
            sys.stderr.write("\n")
            sys.stderr.flush()


def format_duration(seconds: float) -> str:
    seconds_int = int(seconds)
    hours, remainder = divmod(seconds_int, 3600)
    minutes, secs = divmod(remainder, 60)
    if hours:
        return f"{hours:d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def read_inventory(path: Path) -> list[dict[str, object]]:
    intervals: list[dict[str, object]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            if not row.get("start_time") or not row.get("end_time"):
                continue
            if row.get("ffprobe_error"):
                continue
            intervals.append(
                {
                    "path": row["path"],
                    "start": dt.datetime.fromisoformat(row["start_time"]),
                    "end": dt.datetime.fromisoformat(row["end_time"]),
                }
            )
    intervals.sort(key=lambda item: item["start"])
    return intervals


def coverage_parts(start: dt.datetime, end: dt.datetime, intervals: list[dict[str, object]]) -> list[dict[str, object]]:
    parts = []
    for item in intervals:
        item_start = item["start"]
        item_end = item["end"]
        if item_end <= start:
            continue
        if item_start >= end:
            break
        overlap_start = max(start, item_start)
        overlap_end = min(end, item_end)
        if overlap_start < overlap_end:
            parts.append({"start": overlap_start, "end": overlap_end, "path": item["path"]})
    return parts


def classify_coverage(start: dt.datetime, end: dt.datetime, parts: list[dict[str, object]]) -> tuple[str, str, float]:
    if not parts:
        return "no_coverage", "no recording overlaps this interval", 0.0

    cursor = start
    covered_seconds = 0.0
    gaps: list[str] = []
    for part in parts:
        part_start = part["start"]
        part_end = part["end"]
        if part_start > cursor:
            gaps.append(f"{cursor:%Y-%m-%d %H:%M:%S} -> {part_start:%Y-%m-%d %H:%M:%S}")
        if part_end > cursor:
            covered_seconds += (part_end - max(cursor, part_start)).total_seconds()
            cursor = part_end
    if cursor < end:
        gaps.append(f"{cursor:%Y-%m-%d %H:%M:%S} -> {end:%Y-%m-%d %H:%M:%S}")

    total_seconds = (end - start).total_seconds()
    ratio = covered_seconds / total_seconds if total_seconds > 0 else 0.0
    if not gaps:
        status = "fully_covered"
        message = ""
    elif ratio > 0:
        status = "partial_coverage"
        message = "; ".join(gaps)
    else:
        status = "no_coverage"
        message = "no recording overlaps this interval"
    return status, message, ratio


def main() -> int:
    args = parse_args()
    workbook = load_workbook(args.excel, args.password)
    worksheet = workbook[args.sheet] if args.sheet else workbook.active
    intervals = read_inventory(args.inventory)
    ignore_dates = set(args.ignore_dates)

    start_idx = column_index(args.start_col)
    end_idx = column_index(args.end_col)
    label_idx = column_index(args.label_col)
    invalid_idx = None if args.invalid_col.lower() in {"", "none"} else column_index(args.invalid_col)
    print(
        "Resolved columns: "
        f"start={column_letter(start_idx)}, end={column_letter(end_idx)}, "
        f"label={column_letter(label_idx)}, invalid={column_letter(invalid_idx)}"
    )

    rows: list[dict[str, object]] = []
    counts: dict[str, int] = {}
    total_rows = max((getattr(worksheet, "max_row", 0) or 0) - args.first_data_row + 1, 0)
    progress = ProgressReporter(not args.no_progress, total_rows, args.progress_interval_sec)
    processed = 0

    try:
        progress.update(0, force=True)
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_index < args.first_data_row:
                continue
            processed += 1
            progress.update(processed)
            label_raw = row_value(row, label_idx)
            invalid_raw = row_value(row, invalid_idx)
            invalid_signal = normalize_cell_text(invalid_raw) == normalize_cell_text(args.invalid_marker)
            label = sanitize_label(label_raw)
            if not label or label in {"?", "？"}:
                continue
            if invalid_signal and not args.include_invalid_signals:
                counts["invalid_signal"] = counts.get("invalid_signal", 0) + 1
                rows.append(
                    {
                        "row_index": row_index,
                        "label": label,
                        "invalid_signal": str(invalid_raw or ""),
                        "start_time": "",
                        "end_time": "",
                        "duration_sec": "",
                        "coverage_status": "invalid_signal",
                        "coverage_ratio": "0.000000",
                        "message": "marked invalid signal",
                        "source_files": "",
                    }
                )
                continue
            try:
                start = parse_excel_datetime(row_value(row, start_idx))
                end = parse_excel_datetime(row_value(row, end_idx))
                if start is None or end is None or end <= start:
                    status = "invalid_time"
                    message = "missing or invalid start/end time"
                    ratio = 0.0
                    parts: list[dict[str, object]] = []
                elif start.strftime("%Y%m%d") in ignore_dates:
                    status = "ignored_deleted_date"
                    message = "date intentionally removed"
                    ratio = 0.0
                    parts = []
                else:
                    parts = coverage_parts(start, end, intervals)
                    status, message, ratio = classify_coverage(start, end, parts)

                counts[status] = counts.get(status, 0) + 1
                rows.append(
                    {
                        "row_index": row_index,
                        "label": label,
                        "invalid_signal": str(invalid_raw or ""),
                        "start_time": start.isoformat(sep=" ") if start else "",
                        "end_time": end.isoformat(sep=" ") if end else "",
                        "duration_sec": f"{(end - start).total_seconds():.3f}" if start and end and end > start else "",
                        "coverage_status": status,
                        "coverage_ratio": f"{ratio:.6f}",
                        "message": message,
                        "source_files": ";".join(str(part["path"]) for part in parts),
                    }
                )
            except Exception as exc:
                counts["error"] = counts.get("error", 0) + 1
                rows.append(
                    {
                        "row_index": row_index,
                        "label": label,
                        "invalid_signal": str(invalid_raw or ""),
                        "start_time": "",
                        "end_time": "",
                        "duration_sec": "",
                        "coverage_status": "error",
                        "coverage_ratio": "0.000000",
                        "message": str(exc),
                        "source_files": "",
                    }
                )
    finally:
        progress.update(processed, force=True)
        progress.finish()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[
                "row_index",
                "label",
                "invalid_signal",
                "start_time",
                "end_time",
                "duration_sec",
                "coverage_status",
                "coverage_ratio",
                "message",
                "source_files",
            ],
        )
        writer.writeheader()
        writer.writerows(rows)

    if args.summary:
        lines = ["coverage_status_counts:"]
        for key in sorted(counts):
            lines.append(f"  {key}: {counts[key]}")
        invalid_signal_rows = [row for row in rows if row["coverage_status"] == "invalid_signal"]
        uncovered_problem_rows = [
            row
            for row in rows
            if row["coverage_status"] in {"no_coverage", "partial_coverage", "invalid_time", "error"}
        ]
        ready_rows = [row for row in rows if row["coverage_status"] == "fully_covered"]
        lines.append("")
        lines.append("invalid_signal_coverage_check:")
        lines.append(f"  invalid_signal_rows: {len(invalid_signal_rows)}")
        lines.append(f"  fully_covered_rows: {len(ready_rows)}")
        lines.append(f"  remaining_problem_rows_not_marked_invalid: {len(uncovered_problem_rows)}")
        lines.append("")
        lines.append("remaining_problem_row_indices:")
        lines.append(",".join(str(row["row_index"]) for row in uncovered_problem_rows))
        args.summary.parent.mkdir(parents=True, exist_ok=True)
        args.summary.write_text("\n".join(lines) + "\n", encoding="utf-8")

    print(f"Wrote {len(rows)} labelled rows to {args.output}")
    for key in sorted(counts):
        print(f"{key}: {counts[key]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
