#!/usr/bin/env python3
r"""统一生成数据集构建前后的中文检查报告。

默认自动使用 data\annotations\ 下唯一的正式 Excel 标注表。

报告输出到 reports/，包括标注质量、服务器录音库存、行级覆盖率和数据集类别统计。
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import re
import shutil
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path


DEFAULT_ANNOTATION_DIR = Path(r"data\annotations")
SOURCE_RE = re.compile(r"([A-Za-z]+)(\d{8})-(\d{6})\.(\w+)$")
UNRECORDED_CENTER_FREQ_REASON = "中频 5400 不会录到有效信号"

ANNOTATION_STATUS_ZH = {
    "ready_for_extraction": "可提取",
    "invalid_signal": "无效信号",
    "unknown_signal": "未知/未分辨信号",
    "invalid_time": "起止时间无效",
    "unexpected_invalid_value": "无效信号列存在异常值",
}
COVERAGE_STATUS_ZH = {
    "fully_covered": "录音完整覆盖",
    "invalid_signal": "无效信号",
    "invalid_time": "起止时间无效",
    "no_coverage": "无录音覆盖",
    "partial_coverage": "录音部分覆盖",
    "ignored_deleted_date": "已排除日期",
    "error": "检查出错",
}

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="一键生成标注、录音覆盖和数据集统计报告。")
    parser.add_argument("--excel", type=Path, default=None)
    parser.add_argument("--password", default=None)
    parser.add_argument("--sheet", default=None)
    parser.add_argument("--audio-root", type=Path, default=Path(r"D:\AudioRecord"))
    parser.add_argument("--output-dir", type=Path, default=Path("reports"))
    parser.add_argument("--clear-output-dir", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--station", default="LX")
    parser.add_argument("--source-ext", default="m4a")
    parser.add_argument("--start-col", default="F")
    parser.add_argument("--end-col", default="G")
    parser.add_argument("--label-col", default="T")
    parser.add_argument("--center-freq-col", default="B")
    parser.add_argument("--invalid-col", default="N")
    parser.add_argument("--invalid-marker", default="无效信号")
    parser.add_argument("--unrecorded-center-freq", type=float, default=5400.0)
    parser.add_argument("--first-data-row", type=int, default=2)
    parser.add_argument("--ignore-dates", nargs="*", default=[])
    parser.add_argument("--use-existing-inventory", action="store_true")
    parser.add_argument("--ffprobe", default="ffprobe")
    parser.add_argument("--skip-ffprobe", action="store_true")
    parser.add_argument("--no-progress", action="store_true")
    parser.add_argument("--progress-interval-sec", type=float, default=0.2)
    return parser.parse_args()


def discover_annotation_excel(annotation_dir: Path = DEFAULT_ANNOTATION_DIR) -> Path:
    if not annotation_dir.exists():
        raise SystemExit(f"Annotation directory not found: {annotation_dir}")
    candidates = sorted(
        path
        for path in annotation_dir.iterdir()
        if path.is_file()
        and path.suffix.lower() in {".xlsx", ".xlsm"}
        and not path.name.startswith("~$")
    )
    if len(candidates) == 1:
        return candidates[0]
    if not candidates:
        raise SystemExit(f"No annotation Excel file found in {annotation_dir}.")
    names = "\n".join(f"  - {path}" for path in candidates)
    raise SystemExit(
        f"Multiple annotation Excel files found in {annotation_dir}; keep only one or pass --excel explicitly:\n{names}"
    )


def column_index(column: str) -> int:
    result = 0
    for char in column.strip().upper():
        if not ("A" <= char <= "Z"):
            raise ValueError(f"Invalid Excel column: {column}")
        result = result * 26 + (ord(char) - ord("A") + 1)
    return result


def row_value(row: tuple[object, ...], index: int | None):
    if index is None or len(row) < index:
        return None
    return row[index - 1]


def normalize_cell_text(value) -> str:
    return re.sub(r"\s+", "", str(value or "").strip()).lower()


def normalize_label_text(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def parse_float(value) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def is_unrecorded_center_freq(value, target: float) -> bool:
    parsed = parse_float(value)
    return parsed is not None and abs(parsed - target) < 1e-6


def parse_excel_datetime(value) -> dt.datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time.min)
    if isinstance(value, (int, float)):
        from openpyxl.utils.datetime import from_excel

        parsed = from_excel(value)
        if isinstance(parsed, dt.datetime):
            return parsed.replace(tzinfo=None)
        if isinstance(parsed, dt.date):
            return dt.datetime.combine(parsed, dt.time.min)
        return None

    text = str(value).strip()
    if not text:
        return None
    for fmt in (
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y-%m-%d %H:%M",
        "%Y%m%d-%H%M%S",
        "%Y%m%d-%H%M",
    ):
        try:
            return dt.datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise ValueError(f"Unsupported datetime value: {value!r}")


def load_workbook(path: Path, password: str | None):
    from openpyxl import load_workbook

    try:
        return load_workbook(path, data_only=True, read_only=True)
    except Exception:
        if not password:
            raise
        import msoffcrypto

        decrypted = io.BytesIO()
        with path.open("rb") as handle:
            office_file = msoffcrypto.OfficeFile(handle)
            office_file.load_key(password=password)
            office_file.decrypt(decrypted)
        decrypted.seek(0)
        return load_workbook(decrypted, data_only=True, read_only=True)


class ProgressReporter:
    def __init__(self, enabled: bool, total: int | None, interval_sec: float, stage: str) -> None:
        self.enabled = enabled
        self.total = total if total and total > 0 else None
        self.stage = stage
        self.current = 0
        self.bar = None
        if not enabled:
            return
        try:
            from tqdm import tqdm
        except ImportError:
            print(f"{stage}: progress disabled because tqdm is not installed.", flush=True)
            self.enabled = False
            return
        if self.total:
            bar_format = "{desc:<18} [{bar}] {n_fmt}/{total_fmt} {percentage:5.1f}% elapsed={elapsed}"
        else:
            bar_format = "{desc:<18} {n_fmt} elapsed={elapsed}"
        self.bar = tqdm(
            total=self.total,
            desc=stage,
            ascii=True,
            dynamic_ncols=False,
            ncols=80,
            mininterval=max(interval_sec, 0.1),
            file=sys.stdout,
            leave=True,
            bar_format=bar_format,
        )

    def update(self, current: int, force: bool = False) -> None:
        if not self.enabled or self.bar is None:
            return
        delta = current - self.current
        if delta > 0:
            self.bar.update(delta)
            self.current = current
        elif force:
            self.bar.refresh()

    def finish(self) -> None:
        if self.enabled and self.bar is not None:
            self.bar.close()


def format_duration(seconds: float) -> str:
    seconds_int = int(seconds)
    hours, remainder = divmod(seconds_int, 3600)
    minutes, secs = divmod(remainder, 60)
    if hours:
        return f"{hours:d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def write_csv(path: Path, rows: list[dict[str, object]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def safe_clear_report_dir(output_dir: Path) -> None:
    resolved = output_dir.resolve()
    if resolved.anchor == str(resolved):
        raise SystemExit(f"Refusing to clear drive root: {resolved}")
    if resolved.name.lower() != "reports":
        raise SystemExit(f"Refusing to clear {resolved}: report output dir must be named 'reports'.")
    if not resolved.exists():
        return
    if not resolved.is_dir():
        raise SystemExit(f"Refusing to clear non-directory report path: {resolved}")
    for child in resolved.iterdir():
        if child.is_dir():
            shutil.rmtree(child)
        else:
            child.unlink()


def human_dir(args: argparse.Namespace) -> Path:
    return args.output_dir / "human"


def data_dir(args: argparse.Namespace) -> Path:
    return args.output_dir / "data"


def raise_csv_field_limit() -> None:
    limit = sys.maxsize
    while True:
        try:
            csv.field_size_limit(limit)
            return
        except OverflowError:
            limit = int(limit / 10)


def read_annotation_rows(args: argparse.Namespace) -> tuple[list[dict[str, object]], dict[str, int], dict[str, int], str]:
    if not args.no_progress:
        print("Loading Excel annotations...", flush=True)
    workbook = load_workbook(args.excel, args.password)
    worksheet = workbook[args.sheet] if args.sheet else workbook.active
    start_idx = column_index(args.start_col)
    end_idx = column_index(args.end_col)
    label_idx = column_index(args.label_col)
    center_freq_idx = column_index(args.center_freq_col)
    invalid_idx = None if args.invalid_col.lower() in {"", "none"} else column_index(args.invalid_col)
    marker_text = normalize_cell_text(args.invalid_marker)

    total_rows = max((getattr(worksheet, "max_row", 0) or 0) - args.first_data_row + 1, 0)
    progress = ProgressReporter(not args.no_progress, total_rows, args.progress_interval_sec, "Annotation check")
    rows: list[dict[str, object]] = []
    counts: dict[str, int] = {}
    unexpected_invalid_values: dict[str, int] = {}
    processed = 0

    try:
        progress.update(0, force=True)
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_index < args.first_data_row:
                continue
            processed += 1
            progress.update(processed)
            raw_start = row_value(row, start_idx)
            raw_end = row_value(row, end_idx)
            label = str(row_value(row, label_idx) or "").strip()
            center_freq = row_value(row, center_freq_idx)
            invalid_value = str(row_value(row, invalid_idx) or "").strip()
            has_any_key_value = any(value not in (None, "") for value in (raw_start, raw_end, label, center_freq, invalid_value))
            if not has_any_key_value:
                continue

            invalid_norm = normalize_cell_text(invalid_value)
            is_invalid_signal = invalid_norm == marker_text
            unexpected_invalid = bool(invalid_norm and invalid_norm != marker_text)
            if unexpected_invalid:
                unexpected_invalid_values[invalid_value] = unexpected_invalid_values.get(invalid_value, 0) + 1

            try:
                start = parse_excel_datetime(raw_start)
                end = parse_excel_datetime(raw_end)
                time_ok = bool(start and end and end > start)
                time_message = "" if time_ok else "起止时间缺失或结束时间不晚于起始时间"
            except Exception as exc:
                start = None
                end = None
                time_ok = False
                time_message = str(exc)

            label_missing = not label or label in {"?", "？"}
            if unexpected_invalid:
                status = "unexpected_invalid_value"
                message = f"无效信号列存在异常值：{invalid_value!r}"
            elif is_invalid_signal:
                status = "invalid_signal"
                message = "已标注为无效信号"
            elif label_missing:
                status = "unknown_signal"
                message = "推测卫星为空或问号，按未知/未分辨信号跳过"
            elif not time_ok:
                status = "invalid_time"
                message = time_message
            else:
                status = "ready_for_extraction"
                message = ""

            counts[status] = counts.get(status, 0) + 1
            rows.append(
                {
                    "row_index": row_index,
                    "label": label,
                    "center_freq_mhz": center_freq,
                    "invalid_signal": invalid_value,
                    "start_raw": raw_start,
                    "end_raw": raw_end,
                    "start_time": start.isoformat(sep=" ") if start else "",
                    "end_time": end.isoformat(sep=" ") if end else "",
                    "annotation_status": status,
                    "annotation_status_zh": ANNOTATION_STATUS_ZH.get(status, status),
                    "message": message,
                }
            )
    finally:
        progress.update(processed, force=True)
        progress.finish()

    return rows, counts, unexpected_invalid_values, worksheet.title


def parse_source_time(path: Path) -> dt.datetime | None:
    match = SOURCE_RE.match(path.name)
    if not match:
        return None
    _, day, clock, _ = match.groups()
    return dt.datetime.strptime(day + clock, "%Y%m%d%H%M%S")


def ffprobe_duration(ffprobe: str, path: Path) -> tuple[float | None, str]:
    command = [ffprobe, "-v", "error", "-show_entries", "format=duration", "-of", "json", str(path)]
    result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if result.returncode != 0:
        return None, result.stderr.strip()
    try:
        payload = json.loads(result.stdout)
        return float(payload["format"]["duration"]), ""
    except Exception as exc:
        return None, f"Could not parse ffprobe output: {exc}"


def infer_dates(audio_root: Path, station: str) -> list[str]:
    date_dirs = sorted(audio_root.glob(f"{station}[0-9][0-9][0-9][0-9][0-9][0-9][0-9][0-9]"))
    return [path.name.replace(station, "", 1) for path in date_dirs]


def list_audio_files(audio_root: Path, station: str, source_ext: str, dates: list[str]) -> list[Path]:
    files: list[Path] = []
    for day in dates:
        day_dir = audio_root / f"{station}{day}"
        if day_dir.exists():
            files.extend(sorted(day_dir.glob(f"{station}{day}-*.{source_ext.lstrip('.')}")))
    return files


def generate_audio_inventory(args: argparse.Namespace) -> tuple[list[dict[str, object]], list[str], dict[str, int]]:
    output_dir = data_dir(args)
    output_dir.mkdir(parents=True, exist_ok=True)
    dates = infer_dates(args.audio_root, args.station)
    ffprobe_available = shutil.which(args.ffprobe) is not None
    do_ffprobe = (not args.skip_ffprobe) and ffprobe_available

    inventory_rows: list[dict[str, object]] = []
    files_by_day: defaultdict[str, int] = defaultdict(int)
    files = list_audio_files(args.audio_root, args.station, args.source_ext, dates)
    progress = ProgressReporter(not args.no_progress, len(files), args.progress_interval_sec, "Audio inventory")
    progress.update(0, force=True)
    for index, file_path in enumerate(files, start=1):
        start = parse_source_time(file_path)
        duration, probe_error = (None, "")
        if do_ffprobe:
            duration, probe_error = ffprobe_duration(args.ffprobe, file_path)
        end = start + dt.timedelta(seconds=duration) if start and duration is not None else None
        files_by_day[file_path.parent.name] += 1
        stat = file_path.stat()
        inventory_rows.append(
            {
                "path": str(file_path),
                "day_dir": file_path.parent.name,
                "name": file_path.name,
                "start_time": start.isoformat(sep=" ") if start else "",
                "duration_sec": f"{duration:.3f}" if duration is not None else "",
                "end_time": end.isoformat(sep=" ") if end else "",
                "size_bytes": stat.st_size,
                "modified_time": dt.datetime.fromtimestamp(stat.st_mtime).isoformat(sep=" "),
                "ffprobe_error": probe_error,
            }
        )
        progress.update(index)
    progress.update(len(files), force=True)
    progress.finish()

    write_csv(
        output_dir / "audio_inventory.csv",
        inventory_rows,
        ["path", "day_dir", "name", "start_time", "duration_sec", "end_time", "size_bytes", "modified_time", "ffprobe_error"],
    )
    return inventory_rows, dates, dict(files_by_day)


def read_inventory_rows(path: Path) -> list[dict[str, object]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def summarize_inventory_rows(rows: list[dict[str, object]], station: str) -> tuple[list[str], dict[str, int]]:
    dates: set[str] = set()
    files_by_day: Counter[str] = Counter()
    for row in rows:
        day_dir = str(row.get("day_dir") or "")
        if day_dir:
            files_by_day[day_dir] += 1
            if day_dir.startswith(station):
                dates.add(day_dir.replace(station, "", 1))
        elif row.get("start_time"):
            dates.add(str(row["start_time"])[:10].replace("-", ""))
    return sorted(dates), dict(files_by_day)


def read_inventory(path: Path) -> list[dict[str, object]]:
    intervals: list[dict[str, object]] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            if not row.get("start_time") or not row.get("end_time") or row.get("ffprobe_error"):
                continue
            intervals.append(
                {
                    "path": row["path"],
                    "start": dt.datetime.fromisoformat(row["start_time"]),
                    "end": dt.datetime.fromisoformat(row["end_time"]),
                }
            )
    intervals.sort(key=lambda item: item["start"])
    return intervals


def coverage_parts(start: dt.datetime, end: dt.datetime, intervals: list[dict[str, object]]) -> list[dict[str, object]]:
    parts = []
    for item in intervals:
        item_start = item["start"]
        item_end = item["end"]
        if item_end <= start:
            continue
        if item_start >= end:
            break
        overlap_start = max(start, item_start)
        overlap_end = min(end, item_end)
        if overlap_start < overlap_end:
            parts.append({"start": overlap_start, "end": overlap_end, "path": item["path"]})
    return parts


def classify_coverage(start: dt.datetime, end: dt.datetime, parts: list[dict[str, object]]) -> tuple[str, str, float]:
    if not parts:
        return "no_coverage", "没有录音覆盖该时间段", 0.0
    cursor = start
    covered_seconds = 0.0
    gaps: list[str] = []
    for part in parts:
        part_start = part["start"]
        part_end = part["end"]
        if part_start > cursor:
            gaps.append(f"{cursor:%Y-%m-%d %H:%M:%S} -> {part_start:%Y-%m-%d %H:%M:%S}")
        if part_end > cursor:
            covered_seconds += (part_end - max(cursor, part_start)).total_seconds()
            cursor = part_end
    if cursor < end:
        gaps.append(f"{cursor:%Y-%m-%d %H:%M:%S} -> {end:%Y-%m-%d %H:%M:%S}")
    total_seconds = (end - start).total_seconds()
    ratio = covered_seconds / total_seconds if total_seconds > 0 else 0.0
    if not gaps:
        return "fully_covered", "", ratio
    if ratio > 0:
        return "partial_coverage", "; ".join(gaps), ratio
    return "no_coverage", "没有录音覆盖该时间段", ratio


def format_problem_row(row: dict[str, object]) -> str:
    detail_parts = [
        f"行 {row.get('row_index', '')}",
        f"卫星/类别: {row.get('label', '')}",
        f"状态: {row.get('coverage_status_zh', row.get('coverage_status', ''))}",
    ]
    if row.get("center_freq_mhz") not in (None, ""):
        detail_parts.append(f"中频: {row.get('center_freq_mhz')}")
    if row.get("start_time") or row.get("end_time"):
        detail_parts.append(f"时间: {row.get('start_time', '')} -> {row.get('end_time', '')}")
    if row.get("coverage_ratio"):
        detail_parts.append(f"覆盖比例: {row.get('coverage_ratio')}")
    reason = row.get("not_extractable_reason") or row.get("message")
    if reason:
        detail_parts.append(f"原因: {reason}")
    if row.get("source_files"):
        detail_parts.append(f"涉及源文件: {row.get('source_files')}")
    return "；".join(detail_parts)


def generate_row_report(args: argparse.Namespace, annotation_rows: list[dict[str, object]]) -> tuple[list[dict[str, object]], dict[str, int]]:
    inventory_path = data_dir(args) / "audio_inventory.csv"
    if not inventory_path.exists():
        return [], {}
    if not args.no_progress:
        print("Loading Excel annotations for coverage check...", flush=True)
    workbook = load_workbook(args.excel, args.password)
    worksheet = workbook[args.sheet] if args.sheet else workbook.active
    intervals = read_inventory(inventory_path)
    ignore_dates = set(args.ignore_dates)
    start_idx = column_index(args.start_col)
    end_idx = column_index(args.end_col)
    label_idx = column_index(args.label_col)
    center_freq_idx = column_index(args.center_freq_col)
    invalid_idx = None if args.invalid_col.lower() in {"", "none"} else column_index(args.invalid_col)
    marker_text = normalize_cell_text(args.invalid_marker)

    total_rows = max((getattr(worksheet, "max_row", 0) or 0) - args.first_data_row + 1, 0)
    progress = ProgressReporter(not args.no_progress, total_rows, args.progress_interval_sec, "Coverage check")
    rows: list[dict[str, object]] = []
    counts: dict[str, int] = {}
    processed = 0
    try:
        progress.update(0, force=True)
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_index < args.first_data_row:
                continue
            processed += 1
            progress.update(processed)
            raw_start = row_value(row, start_idx)
            raw_end = row_value(row, end_idx)
            label = str(row_value(row, label_idx) or "").strip()
            center_freq_raw = row_value(row, center_freq_idx)
            center_freq_text = "" if center_freq_raw in (None, "") else str(center_freq_raw).strip()
            invalid_raw = row_value(row, invalid_idx)
            has_any_key_value = any(value not in (None, "") for value in (raw_start, raw_end, label, center_freq_raw, invalid_raw))
            if not has_any_key_value:
                continue
            problem_label = normalize_cell_text(invalid_raw) == marker_text
            unrecorded_center_freq = is_unrecorded_center_freq(center_freq_raw, args.unrecorded_center_freq)
            unknown_signal = not label or label in {"?", "？"}
            if unknown_signal:
                status = "unknown_signal"
                counts[status] = counts.get(status, 0) + 1
                rows.append(
                    {
                        "row_index": row_index,
                        "label": label,
                        "center_freq_mhz": center_freq_text,
                        "invalid_signal": str(invalid_raw or ""),
                        "start_time": "",
                        "end_time": "",
                        "duration_sec": "",
                        "annotation_status": "unknown_signal",
                        "annotation_status_zh": ANNOTATION_STATUS_ZH["unknown_signal"],
                        "problem_label": "是" if problem_label else "否",
                        "objective_extractable": "",
                        "coverage_status": "",
                        "coverage_status_zh": "",
                        "final_category": "未参与判断",
                        "label_match_status": "未参与判断",
                        "not_extractable_reason": "未参与判断",
                        "coverage_ratio": "",
                        "message": "推测卫星为空或问号，按未知/未分辨信号跳过",
                        "source_files": "",
                    }
                )
                continue
            try:
                start = parse_excel_datetime(raw_start)
                end = parse_excel_datetime(raw_end)
                if start is None or end is None or end <= start:
                    status = "invalid_time"
                    message = "起止时间缺失或结束时间不晚于起始时间"
                    ratio = 0.0
                    parts: list[dict[str, object]] = []
                elif start.strftime("%Y%m%d") in ignore_dates:
                    status = "ignored_deleted_date"
                    message = "该日期已整体排除"
                    ratio = 0.0
                    parts = []
                else:
                    parts = coverage_parts(start, end, intervals)
                    status, message, ratio = classify_coverage(start, end, parts)
                counts[status] = counts.get(status, 0) + 1
                objective_extractable = status == "fully_covered" and not unrecorded_center_freq
                if status == "invalid_time":
                    not_extractable_reason = "起止时间无效"
                elif unrecorded_center_freq:
                    not_extractable_reason = UNRECORDED_CENTER_FREQ_REASON
                elif status == "no_coverage":
                    not_extractable_reason = "无录音覆盖"
                elif status == "partial_coverage":
                    not_extractable_reason = "录音部分覆盖"
                elif status == "ignored_deleted_date":
                    not_extractable_reason = "已排除日期"
                elif objective_extractable:
                    not_extractable_reason = ""
                else:
                    not_extractable_reason = COVERAGE_STATUS_ZH.get(status, status)
                if objective_extractable and problem_label:
                    label_match_status = "可提取但被标问题信号"
                elif (not objective_extractable) and problem_label:
                    label_match_status = "不可提取且已标问题信号"
                elif (not objective_extractable) and not problem_label:
                    label_match_status = "不可提取但未标问题信号"
                else:
                    label_match_status = "可提取且未标问题信号"
                rows.append(
                    {
                        "row_index": row_index,
                        "label": label,
                        "center_freq_mhz": center_freq_text,
                        "invalid_signal": str(invalid_raw or ""),
                        "start_time": start.isoformat(sep=" ") if start else "",
                        "end_time": end.isoformat(sep=" ") if end else "",
                        "duration_sec": f"{(end - start).total_seconds():.3f}" if start and end and end > start else "",
                        "annotation_status": "invalid_signal" if problem_label else "ready_for_extraction",
                        "annotation_status_zh": "人工标注为问题信号" if problem_label else "未标问题信号",
                        "problem_label": "是" if problem_label else "否",
                        "objective_extractable": "是" if objective_extractable else "否",
                        "coverage_status": status,
                        "coverage_status_zh": COVERAGE_STATUS_ZH.get(status, status),
                        "final_category": "可提取" if objective_extractable else "不可提取",
                        "label_match_status": label_match_status,
                        "not_extractable_reason": not_extractable_reason,
                        "coverage_ratio": f"{ratio:.6f}",
                        "message": message,
                        "source_files": ";".join(str(part["path"]) for part in parts),
                    }
                )
            except Exception as exc:
                status = "error"
                counts[status] = counts.get(status, 0) + 1
                rows.append(
                    {
                        "row_index": row_index,
                        "label": label,
                        "center_freq_mhz": center_freq_text,
                        "invalid_signal": str(invalid_raw or ""),
                        "start_time": "",
                        "end_time": "",
                        "duration_sec": "",
                        "annotation_status": status,
                        "annotation_status_zh": COVERAGE_STATUS_ZH[status],
                        "problem_label": "是" if problem_label else "否",
                        "objective_extractable": "否",
                        "coverage_status": status,
                        "coverage_status_zh": COVERAGE_STATUS_ZH[status],
                        "final_category": "不可提取",
                        "label_match_status": "不可提取且已标问题信号" if problem_label else "不可提取但未标问题信号",
                        "not_extractable_reason": "检查出错",
                        "coverage_ratio": "0.000000",
                        "message": str(exc),
                        "source_files": "",
                    }
                )
    finally:
        progress.update(processed, force=True)
        progress.finish()

    write_csv(
        data_dir(args) / "row_report.csv",
        rows,
        [
            "row_index",
            "label",
            "center_freq_mhz",
            "invalid_signal",
            "start_time",
            "end_time",
            "duration_sec",
            "annotation_status",
            "annotation_status_zh",
            "problem_label",
            "objective_extractable",
            "coverage_status",
            "coverage_status_zh",
            "final_category",
            "label_match_status",
            "not_extractable_reason",
            "coverage_ratio",
            "message",
            "source_files",
        ],
    )
    return rows, counts


def generate_expected_dataset_counts(args: argparse.Namespace, row_report: list[dict[str, object]]) -> list[dict[str, object]]:
    grouped: dict[str, list[dict[str, object]]] = defaultdict(list)
    for row in row_report:
        if row.get("objective_extractable") != "是":
            continue
        label = normalize_label_text(row.get("label"))
        if label:
            grouped[label].append(row)

    report_rows: list[dict[str, object]] = []
    for label in sorted(grouped):
        items = grouped[label]
        row_indices = [int(item["row_index"]) for item in items if str(item.get("row_index") or "").isdigit()]
        start_times = sorted(str(item.get("start_time") or "") for item in items if item.get("start_time"))
        report_rows.append(
            {
                "label": label,
                "expected_count": len(items),
                "first_row_index": min(row_indices) if row_indices else "",
                "last_row_index": max(row_indices) if row_indices else "",
                "first_time": start_times[0] if start_times else "",
                "last_time": start_times[-1] if start_times else "",
            }
        )
    report_rows.sort(key=lambda item: (-int(item["expected_count"]), str(item["label"])))
    write_csv(
        data_dir(args) / "expected_dataset_counts.csv",
        report_rows,
        ["label", "expected_count", "first_row_index", "last_row_index", "first_time", "last_time"],
    )
    return report_rows


def write_expected_dataset_counts_txt(args: argparse.Namespace, expected_rows: list[dict[str, object]]) -> None:
    total = sum(int(row["expected_count"]) for row in expected_rows)
    lines = [
        "预计生成数据集类别统计",
        "",
        f"可生成类别数: {len(expected_rows)}",
        f"可生成音频总数: {total}",
        "",
        "类别明细:",
    ]
    if expected_rows:
        for row in expected_rows:
            lines.append(f"- {row['label']}: {row['expected_count']}")
    else:
        lines.append("无")
    (human_dir(args) / "expected_dataset_counts.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_human_reports(
    args: argparse.Namespace,
    annotation_counts: dict[str, int],
    unexpected_invalid_values: dict[str, int],
    worksheet_title: str,
    row_report: list[dict[str, object]],
    dates: list[str],
    files_by_day: dict[str, int],
    expected_dataset_rows: list[dict[str, object]],
) -> None:
    hdir = human_dir(args)
    hdir.mkdir(parents=True, exist_ok=True)
    category_counts = Counter(row.get("final_category", "") for row in row_report)
    judged_rows = [row for row in row_report if row.get("final_category") != "未参与判断"]
    extractable_rows = [row for row in judged_rows if row.get("objective_extractable") == "是"]
    unextractable_rows = [row for row in judged_rows if row.get("objective_extractable") == "否"]
    problem_labeled_rows = [row for row in judged_rows if row.get("problem_label") == "是"]
    true_problem_labeled_rows = [row for row in judged_rows if row.get("label_match_status") == "不可提取且已标问题信号"]
    missing_problem_label_rows = [row for row in judged_rows if row.get("label_match_status") == "不可提取但未标问题信号"]
    false_problem_label_rows = [row for row in judged_rows if row.get("label_match_status") == "可提取但被标问题信号"]
    unextractable_reason_counts = Counter(
        row.get("not_extractable_reason") or "未分类"
        for row in unextractable_rows
    )
    expected_total = sum(int(row["expected_count"]) for row in expected_dataset_rows)
    coverage_rate = len(true_problem_labeled_rows) / len(unextractable_rows) if unextractable_rows else 1.0
    precision_rate = len(true_problem_labeled_rows) / len(problem_labeled_rows) if problem_labeled_rows else 1.0
    perfect_overlap = not missing_problem_label_rows and not false_problem_label_rows

    lines = [
        "卫星音频数据集分割前检查摘要",
        "",
        "一、结论",
        f"  参与判断总数: {len(judged_rows)}",
        f"  可提取信号: {len(extractable_rows)}",
        f"  不可提取信号: {len(unextractable_rows)}",
        f"  未参与判断: {category_counts.get('未参与判断', 0)}",
        "",
        f"  人工标注为问题信号: {len(problem_labeled_rows)}",
        f"  不可提取且已标问题信号: {len(true_problem_labeled_rows)}",
        f"  不可提取但未标问题信号: {len(missing_problem_label_rows)}",
        f"  可提取但被标问题信号: {len(false_problem_label_rows)}",
        "",
        f"  问题信号标注覆盖率: {coverage_rate:.1%}",
        f"  问题信号标注准确率: {precision_rate:.1%}",
        f"  问题信号标注是否完全重合: {'是' if perfect_overlap else '否'}",
        "",
        "二、不可提取原因",
        f"  {UNRECORDED_CENTER_FREQ_REASON}: {unextractable_reason_counts.get(UNRECORDED_CENTER_FREQ_REASON, 0)}",
        f"  起止时间无效: {unextractable_reason_counts.get('起止时间无效', 0)}",
        f"  无录音覆盖: {unextractable_reason_counts.get('无录音覆盖', 0)}",
        f"  录音部分覆盖: {unextractable_reason_counts.get('录音部分覆盖', 0)}",
        f"  检查出错: {unextractable_reason_counts.get('检查出错', 0)}",
        "",
        "三、预计生成数据集统计",
        f"  可生成类别数: {len(expected_dataset_rows)}",
        f"  可生成音频总数: {expected_total}",
        f"  明细文件: {data_dir(args) / 'expected_dataset_counts.csv'}",
        "",
        "四、Excel 标注检查",
        f"  Excel 文件: {args.excel}",
        f"  工作表: {worksheet_title}",
        f"  列配置: 起始时间={args.start_col}, 结束时间={args.end_col}, 推测卫星={args.label_col}, 无效信号={args.invalid_col}, 中频={args.center_freq_col}",
    ]
    for key in sorted(annotation_counts):
        lines.append(f"  {ANNOTATION_STATUS_ZH.get(key, key)}: {annotation_counts[key]}")
    lines.append("  无效信号列异常值: " + ("无" if not unexpected_invalid_values else ", ".join(f"{k}={v}" for k, v in unexpected_invalid_values.items())))
    lines += [
        "",
        "五、服务器录音库存",
        f"  录音根目录: {args.audio_root}",
        f"  扫描日期: {dates[0] if dates else '无'} 至 {dates[-1] if dates else '无'}",
        f"  日期数量: {len(dates)}",
        f"  录音文件总数: {sum(files_by_day.values())}",
        "",
        "六、输出文件",
        f"  机器明细: {data_dir(args)}",
        f"  人工摘要: {hdir}",
    ]
    (hdir / "summary.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")

    problem_lines = [
        "问题信号标注重合检查明细",
        "",
        "说明: 问题信号 = 客观不可提取的行，包括起止时间无效、录音不完整或没有覆盖，以及中频为 5400 的行。",
        "这里列出客观不可提取但未标问题信号的漏标行，以及客观可提取但被标问题信号的误标行。",
        "",
        "一、不可提取但未标问题信号",
    ]
    if missing_problem_label_rows:
        for row in missing_problem_label_rows:
            problem_lines.append(f"- {format_problem_row(row)}")
    else:
        problem_lines.append("无")
    problem_lines += [
        "",
        "二、可提取但被标问题信号",
    ]
    if false_problem_label_rows:
        for row in false_problem_label_rows:
            problem_lines.append(f"- {format_problem_row(row)}")
    else:
        problem_lines.append("无")
    (hdir / "problem_signals.txt").write_text("\n".join(problem_lines) + "\n", encoding="utf-8")


def main() -> int:
    args = parse_args()
    if args.excel is None:
        args.excel = discover_annotation_excel()
    raise_csv_field_limit()
    if args.clear_output_dir:
        safe_clear_report_dir(args.output_dir)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    human_dir(args).mkdir(parents=True, exist_ok=True)
    data_dir(args).mkdir(parents=True, exist_ok=True)
    annotation_rows, annotation_counts, unexpected_invalid_values, worksheet_title = read_annotation_rows(args)
    inventory_path = data_dir(args) / "audio_inventory.csv"
    if args.use_existing_inventory and inventory_path.exists():
        inventory_rows = read_inventory_rows(inventory_path)
        dates, files_by_day = summarize_inventory_rows(inventory_rows, args.station)
    else:
        inventory_rows, dates, files_by_day = generate_audio_inventory(args)
    row_report, coverage_counts = generate_row_report(args, annotation_rows)
    expected_dataset_rows = generate_expected_dataset_counts(args, row_report)
    write_expected_dataset_counts_txt(args, expected_dataset_rows)
    write_human_reports(
        args,
        annotation_counts,
        unexpected_invalid_values,
        worksheet_title,
        row_report,
        dates,
        files_by_day,
        expected_dataset_rows,
    )
    print(f"Reports written to: {args.output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
